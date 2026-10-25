from pptx import Presentation
from pptx.util import Inches, Pt


def create_presentation(output_path='presentation/SpaceX_Falcon9_Presentation.pptx'):
    prs = Presentation()

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = 'SpaceX Falcon 9 Landing Prediction'
    slide.placeholders[1].text = 'Data Science Portfolio Project\nPredicting first-stage landing success'

    # Slide 2
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = 'Business context'
    body = slide.shapes.placeholders[1].text_frame
    body.text = (
        'SpaceX reduces launch cost by reusing the Falcon 9 first stage.\n'
        'Predicting landing success helps estimate launch economics.\n'
        'This project evaluates mission data and machine learning models.'
    )

    # Slide 3
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = 'Key findings'
    body = slide.shapes.placeholders[1].text_frame
    body.text = (
        '• Flight number and payload mass strongly correlate with landing success.\n'
        '• Launch site and orbit type influence mission outcomes.\n'
        '• Models achieve useful predictive performance with feature engineering.'
    )

    # Slide 4
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = 'Model comparison'
    body = slide.shapes.placeholders[1].text_frame
    body.text = (
        'Logistic Regression\n'
        'Support Vector Machine\n'
        'Decision Tree\n'
        'K-Nearest Neighbors\n'
        'Best performing model is selected based on test accuracy and confusion matrix.'
    )

    # Slide 5
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = 'Conclusion'
    slide.placeholders[1].text = (
        'The project confirms that launch metadata can support a practical first-stage landing prediction pipeline.\n'
        'This is suitable for strong portfolio presentation and business analytics storytelling.'
    )

    prs.save(output_path)
    print(f'PowerPoint saved to: {output_path}')


if __name__ == '__main__':
    create_presentation()
